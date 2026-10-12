"""Bounded, local text extraction. Never executes document content."""
import base64,io,pathlib,zipfile
MAX_BYTES=5*1024*1024
MAX_CHARS=6000
def extract_document(data):
    name=pathlib.PurePath(str(data.get('name','document'))).name[:150]
    suffix=pathlib.PurePath(name).suffix.lower()
    if suffix not in {'.txt','.md','.pdf','.docx'}: raise ValueError('Supported files: TXT, Markdown, text-based PDF and DOCX.')
    try: raw=base64.b64decode(data['base64'],validate=True)
    except Exception: raise ValueError('Invalid file data.')
    if not raw or len(raw)>MAX_BYTES: raise ValueError('Choose a non-empty file no larger than 5 MB.')
    if suffix in {'.txt','.md'}:
        try: text=raw.decode('utf-16' if raw[:2] in {b'\xff\xfe',b'\xfe\xff'} else 'utf-8-sig')
        except UnicodeError: raise ValueError('Text files must use UTF-8 or UTF-16 encoding.')
        if '\x00' in text: raise ValueError('This appears to be a binary file, not readable text.')
    elif suffix=='.pdf':
        from pypdf import PdfReader
        try:
            reader=PdfReader(io.BytesIO(raw))
            if reader.is_encrypted: raise ValueError('Password-protected PDFs are unsupported.')
            if len(reader.pages)>100: raise ValueError('Use a PDF with at most 100 pages.')
            parts=[]
            for page in reader.pages:
                parts.append(page.extract_text() or '')
                if sum(map(len,parts))>MAX_CHARS: raise ValueError('The extracted document is too long. Use an excerpt of at most 6,000 characters.')
            text='\n'.join(parts)
        except ValueError: raise
        except Exception: raise ValueError('Unable to read this PDF. Choose a valid text-based PDF.')
    else:
        from docx import Document
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                info=archive.infolist()
                if len(info)>500 or sum(x.file_size for x in info)>20*1024*1024: raise ValueError('DOCX archive is too large to extract safely.')
                if any(x.flag_bits&1 for x in info): raise ValueError('Encrypted DOCX files are unsupported.')
            document=Document(io.BytesIO(raw))
            parts=[p.text for p in document.paragraphs]
            parts.extend(' | '.join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
            text='\n'.join(parts)
        except ValueError: raise
        except Exception: raise ValueError('Unable to read this DOCX. Choose a valid Word document.')
    text=text.strip()
    if not text: raise ValueError('No readable text found. Scanned PDFs and image-only documents need OCR, which this app does not support.')
    if len(text)>MAX_CHARS: raise ValueError('The extracted document is too long. Use an excerpt of at most 6,000 characters.')
    return {'name':name,'text':text,'characters':len(text),'bytes':len(raw),'note':'Text extracted locally. Images, OCR, headers, footers and complex layout are not included.'}
